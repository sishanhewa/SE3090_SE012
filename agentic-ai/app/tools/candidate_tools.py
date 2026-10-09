"""Candidate Analysis Agent tools — read-only data retrieval from backend API."""
from app.clients.backend_url import backend_api_url
import httpx
import structlog
from io import BytesIO
from pathlib import Path
import re
from pypdf import PdfReader
from docx import Document
from typing import Any, Optional

logger = structlog.get_logger()

# Base URL for the ASP.NET backend API

BACKEND_BASE_URL = backend_api_url()


async def get_application(
    application_id: str, auth_token: Optional[str] = None
) -> dict[str, Any]:
    """Fetch application details including the candidate_profile_id."""
    async with httpx.AsyncClient(timeout=30.0) as client:
        headers = _auth_headers(auth_token)
        response = await client.get(
            f"{BACKEND_BASE_URL}/applications/{application_id}", headers=headers
        )
        response.raise_for_status()
        app_data = response.json()
        
        logger.info("tool_get_application", application_id=application_id, status="success")
        return {
            "application_id": application_id,
            "job_id": app_data.get("jobId", ""),
            "candidate_profile_id": app_data.get("candidateProfileId", ""),
            "status": app_data.get("status", 0),
            "cover_letter": app_data.get("coverLetter", ""),
            "candidate_name": app_data.get("candidateName", ""),
            "resume_document_id": app_data.get("resumeDocumentId"),
        }


async def get_application_resume(
    application_id: str, auth_token: Optional[str] = None
) -> dict[str, str]:
    """Read only the CV attached to this application, through the protected API."""
    async with httpx.AsyncClient(timeout=30.0) as client:
        response = await client.get(
            f"{BACKEND_BASE_URL}/applications/{application_id}/resume",
            headers=_auth_headers(auth_token),
        )
        response.raise_for_status()
        content = response.content
        if not content or len(content) > 10_000_000:
            raise ValueError("The attached CV is empty or exceeds 10 MB")

        disposition = response.headers.get("content-disposition", "")
        match = re.search(r'filename="?([^";]+)', disposition, re.IGNORECASE)
        file_name = Path(match.group(1)).name if match else "CV"
        suffix = Path(file_name).suffix.lower()
        if not suffix:
            content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
            suffix = ".pdf" if content_type == "application/pdf" else ".docx" if "wordprocessingml" in content_type else ""
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages[:30])
        elif suffix == ".docx":
            document = Document(BytesIO(content))
            text = "\n".join(p.text for p in document.paragraphs)
            text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        else:
            raise ValueError("CV screening supports PDF and DOCX files only")

        text = " ".join(text.split())[:20_000]
        if len(text) < 30:
            raise ValueError("No readable text found in the CV; upload a text-based PDF or DOCX")
        logger.info("tool_get_application_resume", application_id=application_id, characters=len(text))
        return {"file_name": file_name, "text": text}

async def get_job_requirements(
    job_id: str, auth_token: Optional[str] = None
) -> dict[str, Any]:
    """
    Fetch job requirements including skills, experience, and qualifications.
    
    Allowed tool for: CandidateAnalysisAgent
    """
    async with httpx.AsyncClient(timeout=30.0) as client:
        headers = _auth_headers(auth_token)
        response = await client.get(
            f"{BACKEND_BASE_URL}/jobs/{job_id}", headers=headers
        )
        response.raise_for_status()
        job_data = response.json()

        logger.info("tool_get_job_requirements", job_id=job_id, status="success")
        return {
            "job_id": job_id,
            "title": job_data.get("title", ""),
            "description": job_data.get("description", ""),
            "minimum_experience": job_data.get("minimumExperience", 0),
            "vacancy_count": job_data.get("vacancyCount", 0),
            "employment_type": job_data.get("employmentType", ""),
            "skill_requirements": job_data.get("skillRequirements", []),
            "requirements": job_data.get("requirements", []),
        }


async def get_candidate_profile(
    candidate_profile_id: str, auth_token: Optional[str] = None
) -> dict[str, Any]:
    """
    Fetch candidate profile with summary and contact info.
    
    Allowed tool for: CandidateAnalysisAgent
    """
    async with httpx.AsyncClient(timeout=30.0) as client:
        headers = _auth_headers(auth_token)
        response = await client.get(
            f"{BACKEND_BASE_URL}/candidateprofiles/{candidate_profile_id}",
            headers=headers,
        )
        response.raise_for_status()
        profile = response.json()

        logger.info(
            "tool_get_candidate_profile",
            candidate_id=candidate_profile_id,
            status="success",
        )
        return {
            "candidate_profile_id": candidate_profile_id,
            "summary": profile.get("summary", ""),
            "phone": profile.get("phone", ""),
            "linkedin_url": profile.get("linkedInUrl", ""),
            "portfolio_url": profile.get("portfolioUrl", ""),
        }


async def get_candidate_skills(
    candidate_profile_id: str, auth_token: Optional[str] = None
) -> list[dict[str, Any]]:
    """
    Fetch candidate skills with proficiency levels and experience.
    
    Allowed tool for: CandidateAnalysisAgent
    """
    async with httpx.AsyncClient(timeout=30.0) as client:
        headers = _auth_headers(auth_token)
        response = await client.get(
            f"{BACKEND_BASE_URL}/candidateprofiles/{candidate_profile_id}",
            headers=headers,
        )
        response.raise_for_status()
        profile = response.json()

        skills = profile.get("skills", [])
        logger.info(
            "tool_get_candidate_skills",
            candidate_id=candidate_profile_id,
            skill_count=len(skills),
        )
        return skills


async def get_candidate_experience(
    candidate_profile_id: str, auth_token: Optional[str] = None
) -> list[dict[str, Any]]:
    """
    Fetch candidate work experience history.
    
    Allowed tool for: CandidateAnalysisAgent
    """
    async with httpx.AsyncClient(timeout=30.0) as client:
        headers = _auth_headers(auth_token)
        response = await client.get(
            f"{BACKEND_BASE_URL}/candidateprofiles/{candidate_profile_id}",
            headers=headers,
        )
        response.raise_for_status()
        profile = response.json()

        experience = profile.get("experience", [])
        logger.info(
            "tool_get_candidate_experience",
            candidate_id=candidate_profile_id,
            experience_count=len(experience),
        )
        return experience


async def get_candidate_education(
    candidate_profile_id: str, auth_token: Optional[str] = None
) -> list[dict[str, Any]]:
    """
    Fetch candidate education history.
    
    Allowed tool for: CandidateAnalysisAgent
    """
    async with httpx.AsyncClient(timeout=30.0) as client:
        headers = _auth_headers(auth_token)
        response = await client.get(
            f"{BACKEND_BASE_URL}/candidateprofiles/{candidate_profile_id}",
            headers=headers,
        )
        response.raise_for_status()
        profile = response.json()

        education = profile.get("education", [])
        logger.info(
            "tool_get_candidate_education",
            candidate_id=candidate_profile_id,
            education_count=len(education),
        )
        return education


async def get_application_documents(
    candidate_profile_id: str, auth_token: Optional[str] = None
) -> list[dict[str, Any]]:
    """
    Fetch candidate uploaded documents (CV, certificates, etc.).
    
    Allowed tool for: CandidateAnalysisAgent
    """
    async with httpx.AsyncClient(timeout=30.0) as client:
        headers = _auth_headers(auth_token)
        response = await client.get(
            f"{BACKEND_BASE_URL}/candidateprofiles/{candidate_profile_id}",
            headers=headers,
        )
        response.raise_for_status()
        profile = response.json()

        documents = profile.get("documents", [])
        logger.info(
            "tool_get_application_documents",
            candidate_id=candidate_profile_id,
            document_count=len(documents),
        )
        return documents


def _auth_headers(token: Optional[str] = None) -> dict[str, str]:
    """Build authorization headers."""
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    return headers
